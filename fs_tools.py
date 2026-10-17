"""Structured file-system tools for resume documents (PDF, TXT, DOCX).

The functions return JSON-serializable dictionaries/lists and are designed to
be exposed as tools to an LLM. File operations are intentionally local-only.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any
import re

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def _metadata(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "filepath": str(path.resolve()),
        "extension": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(
            stat.st_mtime, tz=timezone.utc
        ).isoformat(),
    }


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support requires pypdf. Install requirements.txt.") from exc
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx. Install requirements.txt.") from exc
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Include table cells, where resume content is commonly stored.
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs).strip()
    raise ValueError(f"Unsupported file extension '{path.suffix}'. Supported: PDF, TXT, DOCX.")


def read_file(filepath: str) -> dict[str, Any]:
    """Read and extract text from a PDF, TXT, or DOCX file."""
    path = Path(filepath).expanduser()
    try:
        if not path.exists():
            return {"success": False, "error": "file_not_found", "message": f"File not found: {filepath}"}
        if not path.is_file():
            return {"success": False, "error": "not_a_file", "message": f"Not a file: {filepath}"}
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            return {"success": False, "error": "unsupported_format",
                    "message": f"Unsupported format '{path.suffix}'. Supported: PDF, TXT, DOCX."}
        content = _extract_text(path)
        return {"success": True, "content": content, "metadata": _metadata(path),
                "character_count": len(content)}
    except Exception as exc:
        return {"success": False, "error": "read_error", "message": str(exc), "filepath": str(path)}


def list_files(directory: str, extension: str | None = None) -> list[dict[str, Any]]:
    """List files directly inside a directory, optionally filtered by extension."""
    path = Path(directory).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")
    if not path.is_dir():
        raise NotADirectoryError(f"Not a directory: {directory}")
    normalized = None
    if extension:
        normalized = extension.lower().strip()
        if not normalized.startswith("."):
            normalized = "." + normalized
    results = []
    for item in sorted(path.iterdir(), key=lambda p: p.name.lower()):
        if item.is_file() and (normalized is None or item.suffix.lower() == normalized):
            try:
                results.append(_metadata(item))
            except OSError:
                continue
    return results


def write_file(filepath: str, content: str) -> dict[str, Any]:
    """Write UTF-8 text, creating parent directories when required."""
    path = Path(filepath).expanduser()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return {"success": True, "filepath": str(path.resolve()),
                "bytes_written": path.stat().st_size, "message": "File written successfully."}
    except Exception as exc:
        return {"success": False, "error": "write_error", "message": str(exc), "filepath": str(path)}


def search_in_file(filepath: str, keyword: str) -> dict[str, Any]:
    """Case-insensitive search with surrounding text context."""
    if not keyword or not keyword.strip():
        return {"success": False, "error": "invalid_keyword", "message": "Keyword must not be empty."}
    result = read_file(filepath)
    if not result.get("success"):
        return result
    content = result["content"]
    matches = []
    for match in re.finditer(re.escape(keyword), content, flags=re.IGNORECASE):
        start, end = match.span()
        context_start, context_end = max(0, start - 70), min(len(content), end + 70)
        matches.append({
            "matched_text": content[start:end],
            "start_index": start,
            "end_index": end,
            "context": content[context_start:context_end].replace("\n", " "),
        })
    return {"success": True, "filepath": str(Path(filepath).expanduser().resolve()),
            "keyword": keyword, "match_count": len(matches), "matches": matches,
            "metadata": result["metadata"]}
